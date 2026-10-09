# -*- coding: utf-8 -*-
"""Разбор таблиц в данные для показа настоящими ячейками.

Раньше xlsx открывался так же, как документ: сервер превращал его в PDF, и
человек видел картинку таблицы — без выделения ячеек, без прокрутки по
столбцам, с разрывами страниц посреди данных. Здесь книга разбирается в
обычные данные, а сетку рисует уже браузер.

Отдаётся только то, что нужно для показа: значения как их видит человек,
ширины столбцов, объединения, закрёпленные области и минимум оформления
(жирность, курсив, выравнивание, цвет заливки и текста).
"""
import csv
import gzip
import io
import logging
import zipfile
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

SHEET_EXTENSIONS = frozenset({".xlsx", ".xlsm", ".xltx", ".xltm", ".csv", ".tsv"})

# Пределы: книга может быть огромной, а браузеру её ещё рисовать
MAX_ROWS = 3000
MAX_COLS = 200
MAX_SHEETS = 30
MAX_TEXT = 500                      # длина значения одной ячейки


def is_sheet_name(name: str) -> bool:
    return Path(str(name or "")).suffix.lower() in SHEET_EXTENSIONS


def _open(stored: Path) -> io.BytesIO:
    """Файл в хранилище может лежать сжатым — читаем в память как есть"""
    if stored.suffix.lower() == ".gz":
        with gzip.open(stored, "rb") as f:
            return io.BytesIO(f.read())
    return io.BytesIO(stored.read_bytes())


def _letters(index: int) -> str:
    """Номер столбца → его буква: 1 → A, 27 → AA"""
    out = ""
    while index > 0:
        index, rest = divmod(index - 1, 26)
        out = chr(65 + rest) + out
    return out


def _text(value, number_format: str = "") -> str:
    """Значение так, как его видит человек в редакторе"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "ИСТИНА" if value else "ЛОЖЬ"
    if isinstance(value, datetime):
        fmt = (number_format or "").lower()
        if value.hour or value.minute or value.second or "h" in fmt:
            return value.strftime("%d.%m.%Y %H:%M")
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, timedelta):
        total = int(value.total_seconds())
        return "%d:%02d:%02d" % (total // 3600, total // 60 % 60, total % 60)
    if isinstance(value, float):
        # 3.0 показываем как 3, а не как 3.0 — как это делает редактор
        if value == int(value) and abs(value) < 1e15:
            return str(int(value))
        return ("%.10f" % value).rstrip("0").rstrip(".")
    return str(value)[:MAX_TEXT]


def _rgb(color) -> Optional[str]:
    """Цвет openpyxl → #RRGGBB. Тема и индексы палитры пропускаются."""
    if color is None or getattr(color, "type", None) != "rgb":
        return None
    value = getattr(color, "rgb", None)
    if not isinstance(value, str) or len(value) != 8:
        return None
    if value[:2] == "00":                    # полностью прозрачный
        return None
    return "#" + value[2:]


def _cell_style(cell) -> dict:
    """Оформление ячейки — только то, что видно в таблице"""
    out = {}
    font, fill, align = cell.font, cell.fill, cell.alignment
    if font is not None:
        if font.bold:
            out["b"] = 1
        if font.italic:
            out["i"] = 1
        color = _rgb(font.color)
        if color and color != "#000000":
            out["fg"] = color
    if fill is not None and getattr(fill, "patternType", None) == "solid":
        color = _rgb(fill.fgColor)
        if color and color != "#FFFFFF":
            out["bg"] = color
    if align is not None:
        if align.horizontal in ("center", "right"):
            out["a"] = align.horizontal
        if align.wrapText:
            out["w"] = 1
    return out


def _csv_sheet(stored: Path, name: str) -> dict:
    """CSV и TSV: разделитель и кодировку определяем сами"""
    raw = _open(stored).read()
    text = None
    for encoding in ("utf-8-sig", "cp1251", "utf-16", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if text is None:
        text = raw.decode("utf-8", "replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if name.lower().endswith(".tsv") else (";" if sample.count(";") > sample.count(",") else ",")
    rows = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        rows.append([{"v": str(v)[:MAX_TEXT]} for v in row[:MAX_COLS]])
        if len(rows) >= MAX_ROWS:
            break
    width = max((len(r) for r in rows), default=0)
    for row in rows:
        row.extend({} for _ in range(width - len(row)))
    return {"name": Path(name).stem[:31] or "Данные", "rows": rows,
            "cols": [{} for _ in range(width)], "merges": [], "freeze": None}


def _book_sheets(stored: Path) -> list:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    book = load_workbook(_open(stored), data_only=True, read_only=False)
    sheets = []
    for ws in book.worksheets[:MAX_SHEETS]:
        if ws.sheet_state != "visible":
            continue
        height = min(ws.max_row or 0, MAX_ROWS)
        width = min(ws.max_column or 0, MAX_COLS)
        rows = []
        for r in range(1, height + 1):
            row = []
            for c in range(1, width + 1):
                cell = ws.cell(row=r, column=c)
                item = _cell_style(cell)
                text = _text(cell.value, cell.number_format)
                if text:
                    item["v"] = text
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        item["n"] = 1          # число — прижимаем вправо
                row.append(item)
            rows.append(row)

        cols = []
        for c in range(1, width + 1):
            dim = ws.column_dimensions.get(get_column_letter(c))
            item = {}
            if dim is not None:
                if dim.hidden:
                    item["h"] = 1
                elif dim.width:
                    item["w"] = round(float(dim.width), 1)
            cols.append(item)

        merges = [{"r": m.min_row - 1, "c": m.min_col - 1,
                   "rs": m.max_row - m.min_row + 1, "cs": m.max_col - m.min_col + 1}
                  for m in ws.merged_cells.ranges
                  if m.min_row <= height and m.min_col <= width]

        freeze = None
        if ws.freeze_panes and ws.freeze_panes != "A1":
            try:
                anchor = ws[ws.freeze_panes]
                freeze = {"r": anchor.row - 1, "c": anchor.column - 1}
            except Exception:
                freeze = None

        sheets.append({"name": ws.title, "rows": rows, "cols": cols,
                       "merges": merges, "freeze": freeze,
                       "truncated": bool((ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLS)})
    book.close()
    return sheets


def sheet_json(stored: Path, display_name: str) -> Optional[dict]:
    """Книга в виде данных для браузера. None — если это не таблица.

    Ошибку разбора не превращаем в отказ сервера: клиент в этом случае
    просто откроет файл как раньше или предложит скачать.
    """
    if not is_sheet_name(display_name):
        return None
    if not stored.is_file():
        return None
    suffix = Path(display_name).suffix.lower()
    try:
        if suffix in (".csv", ".tsv"):
            sheets = [_csv_sheet(stored, display_name)]
        else:
            sheets = _book_sheets(stored)
    except zipfile.BadZipFile:
        logger.info("Таблица «%s»: файл повреждён или это не xlsx", display_name)
        return None
    except ImportError:
        logger.warning("Показ таблиц недоступен: не установлен openpyxl "
                       "(pip install openpyxl)")
        return None
    except Exception as e:
        logger.warning("Таблица «%s» не разобралась: %s", display_name, e)
        return None
    if not sheets:
        return None
    return {"filename": display_name, "sheets": sheets,
            "letters": [_letters(i + 1) for i in range(max(len(s["cols"]) for s in sheets))]}
